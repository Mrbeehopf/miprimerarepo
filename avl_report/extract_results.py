"""Read the computed (LibreOffice-recalculated) values back out of the
Excel workbook into the per-channel JSON structure expected by
build_pdf.py (`chandata.json` in the reference scripts).
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import List

import openpyxl

from .models import CalibrationData
from .channels import DATA_START_ROW, channel_specs


def extract_channel_results(xlsx_path: Path, data: CalibrationData) -> List[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    results = []
    for spec in channel_specs():
        ws = wb[spec.sheet]
        npts = spec.npts(data)
        d1 = DATA_START_ROW
        rows = []
        for i in range(npts):
            r = d1 + i
            ref = ws.cell(r, 2).value
            samp = ws.cell(r, 3).value
            dv = ws.cell(r, 4).value
            dp = ws.cell(r, 5).value
            if ref is None or samp is None or dv is None or dp is None:
                raise ValueError(
                    f"Sheet '{spec.sheet}' row {r}: missing computed value "
                    "(workbook not recalculated?)"
                )
            rows.append((float(ref), float(samp), float(dv), float(dp)))

        def cell(coord):
            v = ws[coord].value
            if v is None:
                raise ValueError(f"Sheet '{spec.sheet}' cell {coord}: missing computed value")
            return v

        results.append({
            "sheet": spec.sheet,
            "range": spec.range_label,
            "a1": float(cell("D21")),
            "a1lim": str(cell("E21")),
            "ic": float(cell("D22")),
            "iclim": str(cell("E22")),
            "see": float(cell("D23")),
            "seelim": str(cell("E23")),
            "r2": float(cell("D24")),
            "r2lim": str(cell("E24")),
            "state": str(cell("D25")),
            "rows": rows,
        })
    return results


def save_results(results: List[dict], out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    return out_path
