"""Build the Linearity Verification Excel workbook (14 sheets) from a
CalibrationData object.

Ported from the validated reference script `reference_build_xlsx.py`; the
exact calculation logic, cell layout and styling are kept unchanged. The
only difference is that data comes from a `CalibrationData` instance
(extracted per certificate) instead of a hardcoded `data` module.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as gcl

from .models import CalibrationData
from .channels import DATA_START_ROW, channel_specs

BOLD = Font(name="Calibri", bold=True, size=11)
NORM = Font(name="Calibri", size=11)
TITLE = Font(name="Calibri", bold=False, size=20)
SECTION = PatternFill("solid", fgColor="E7E6E6")
TBLHDR = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFFF00")
GREENF = PatternFill("solid", fgColor="92D050")
REDF = PatternFill("solid", fgColor="FF0000")

REF_RANGE = "References!$B$2:$G$16"
PROTO = "X-Meter CAL PROTOCOL"


def _result_cf(ws, coord):
    ws.conditional_formatting.add(coord, CellIsRule(operator="equal", formula=['"PASSED"'], fill=GREENF))
    ws.conditional_formatting.add(coord, CellIsRule(operator="equal", formula=['"FAILED"'], fill=REDF))


def _snum(cell, fmt="0.00000"):
    cell.font = NORM
    cell.number_format = fmt


def _sheet_cal_info(wb, info: dict):
    ws = wb.active
    ws.title = "Cal-Info"
    rows = [
        ("AVL Calibration Certificate", None), (None, None),
        ("Calibration Date", info["Calibration Date"]), ("Report ID", info["Report ID"]), (None, None),
        ("Object", None), ("Device", info["Device"]), ("Serial No", info["Serial No"]), (None, None),
        ("Traceability / Calibration Standards", None),
        ("DVM Manufacturer", info["DVM Manufacturer"]), ("DVM Type", info["DVM Type"]),
        ("DVM Gage No", info["DVM Gage No"]), ("DVM Calibration Date", info["DVM Calibration Date"]), (None, None),
        ("Measurement Conditions", None),
        ("Ambient Temperature", info["Ambient Temperature"]), ("Relative Humidity", info["Relative Humidity"]), (None, None),
        ("Notes", None), ("(1) Measurement Uncertainty U95", None), ("(2) Nominal Calibration Voltage", None),
    ]
    for i, (a, b) in enumerate(rows, 1):
        if a is not None:
            ws.cell(i, 1, a).font = BOLD
        if b is not None:
            ws.cell(i, 2, b).font = NORM
    ws.cell(1, 1).font = Font(name="Calibri", bold=True, size=20)
    for hr in (6, 10, 16, 20):
        ws.cell(hr, 1).fill = SECTION
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40


def _sheet_references(wb):
    ws = wb.create_sheet("References")
    refhdr = ["System", "Units", "|xmin(a1-1)+a0|", "a1", "SEE", "r2"]
    data_ref = [
        ("Batch sampler flow rates", "[lpm]", 0.01, 0.02, 0.02, 0.99),
        ("Current", "[Amps]", 0.01, 0.02, 0.02, 0.99),
        ("Dewpoint (Intake Air, PM)", "[°C]", 0.005, 0.01, 0.005, 0.998),
        ("Dewpoint (Other)", "[°C]", 0.01, 0.01, 0.01, 0.998),
        ("Diluted exhaust flow rate", "[m3/min]", 0.01, 0.02, 0.02, 0.99),
        ("Dilution air flow rate", "[m3/min]", 0.01, 0.02, 0.02, 0.99),
        ("Electrical power", "[Watts]", 0.01, 0.02, 0.02, 0.99),
        ("Gas analyzers for field testing", "[ppm]", 0.01, 0.01, 0.01, 0.998),
        ("Gas analyzers for laboratory testing", "[ppm]", 0.005, 0.01, 0.01, 0.998),
        ("Intake-air flow rate", "[kg/s]", 0.01, 0.02, 0.02, 0.99),
        ("PM balance", "[µg]", 0.01, 0.01, 0.01, 0.998),
        ("Pressure", "[hPa]", 0.01, 0.01, 0.01, 0.998),
        ("Raw Exhaust Flow Rate", "[m3/min]", 0.01, 0.02, 0.02, 0.99),
        ("Temperature", "[°C]", 0.01, 0.01, 0.01, 0.998),
        ("Voltage", "[V]", 0.01, 0.02, 0.02, 0.99),
    ]
    for j, h in enumerate(refhdr, 2):
        hc = ws.cell(1, j, h)
        hc.font = BOLD
        hc.fill = TBLHDR
    for i, row in enumerate(data_ref, 2):
        for j, val in enumerate(row, 2):
            c = ws.cell(i, j, val)
            c.font = NORM
            if j >= 4:
                c.number_format = "0.000"
    for col, w in {"B": 38, "C": 12, "D": 16, "E": 8, "F": 8, "G": 8}.items():
        ws.column_dimensions[col].width = w


def _sheet_protocol(wb, data: CalibrationData):
    ws = wb.create_sheet(PROTO)
    r10, r20, r75 = data.r10, data.r20, data.r75

    c = ws.cell(3, 2, "Range +/- 10V")
    c.font = BOLD
    c.fill = SECTION
    ws.cell(4, 2, f"Accuracy Specification: {r10.accuracy_spec}").font = NORM
    for k, v in enumerate(r10.voltages):
        _snum(ws.cell(6, 3 + k), "0.0")
        ws.cell(6, 3 + k).value = v
    ws.cell(6, 2, "Calibration Voltage (2) [V]").font = BOLD
    for ch in range(1, 9):
        ws.cell(7 + ch, 2, f"Error Channel{ch}").font = NORM
        for k, e in enumerate(r10.errors[ch]):
            _snum(ws.cell(7 + ch, 3 + k))
            ws.cell(7 + ch, 3 + k).value = e
    for ch in range(1, 9):
        ws.cell(17 + ch, 2, f"Reading Channel{ch}").font = NORM
        for k in range(len(r10.voltages)):
            L = gcl(3 + k)
            ws.cell(17 + ch, 3 + k).value = f"={L}$6+{L}{7 + ch}"
            _snum(ws.cell(17 + ch, 3 + k))
    ws["L18"] = ArrayFormula("L18:L25", "=TRANSPOSE(C6:J6)")
    for k in range(len(r10.voltages)):
        _snum(ws.cell(18 + k, 12), "0.0")
    for ch in range(1, 9):
        col = 12 + ch
        L = gcl(col)
        ws[f"{L}18"] = ArrayFormula(f"{L}18:{L}25", f"=TRANSPOSE(C{17 + ch}:J{17 + ch})")
        for k in range(len(r10.voltages)):
            _snum(ws.cell(18 + k, col))

    c = ws.cell(28, 2, "Range +/- 20V")
    c.font = BOLD
    c.fill = SECTION
    ws.cell(29, 2, f"Accuracy Specification: {r20.accuracy_spec}").font = NORM
    for k, v in enumerate(r20.voltages):
        _snum(ws.cell(31, 3 + k), "0.0")
        ws.cell(31, 3 + k).value = v
    ws.cell(31, 2, "Calibration Voltage (2) [V]").font = BOLD
    ws.cell(32, 2, "Error Channel9 [V]").font = NORM
    for k, e in enumerate(r20.errors[9]):
        _snum(ws.cell(32, 3 + k))
        ws.cell(32, 3 + k).value = e
    ws["L30"] = ArrayFormula("L30:L36", "=TRANSPOSE(C31:I31)")
    ws["M30"] = ArrayFormula("M30:M36", "=TRANSPOSE(C32:I32)")
    for k in range(len(r20.voltages)):
        _snum(ws.cell(30 + k, 12), "0.0")
        _snum(ws.cell(30 + k, 13))
        ws.cell(30 + k, 14).value = f"=L{30 + k}+M{30 + k}"
        _snum(ws.cell(30 + k, 14))

    c = ws.cell(40, 2, "Range +/- 75V")
    c.font = BOLD
    c.fill = SECTION
    ws.cell(41, 2, f"Accuracy Specification: {r75.accuracy_spec}").font = NORM
    for k, v in enumerate(r75.voltages):
        _snum(ws.cell(42, 3 + k), "0.0")
        ws.cell(42, 3 + k).value = v
    ws.cell(42, 2, "Calibration Voltage (2)").font = BOLD
    ws.cell(43, 2, "Error Channel10").font = NORM
    for k, e in enumerate(r75.errors[10]):
        _snum(ws.cell(43, 3 + k))
        ws.cell(43, 3 + k).value = e
    ws["L41"] = ArrayFormula("L41:L47", "=TRANSPOSE(C42:I42)")
    ws["M41"] = ArrayFormula("M41:M47", "=TRANSPOSE(C43:I43)")
    for k in range(len(r75.voltages)):
        _snum(ws.cell(41 + k, 12), "0.0")
        _snum(ws.cell(41 + k, 13))
        ws.cell(41 + k, 14).value = f"=L{41 + k}+M{41 + k}"
        _snum(ws.cell(41 + k, 14))

    ws.column_dimensions["B"].width = 26
    for col in "CDEFGHIJ":
        ws.column_dimensions[col].width = 11
    for col in "LMNOPQRST":
        ws.column_dimensions[col].width = 12


def _make_channel(wb, sheetname, inst_range, npts, vsc, info: dict):
    ws = wb.create_sheet(sheetname)
    tc = ws.cell(1, 1, '=CONCATENATE("Linearity Verification - ",B3)')
    tc.font = TITLE
    meta = [
        (2, "Test Data", None), (3, "Measurement Type:", "Voltage"), (4, "Instrument:", "AVL X-Meter"),
        (5, "Serial Number:", info["Serial No"]), (6, "Instrument Range:", inst_range),
        (7, "Order / Device:", "SEMA / X-Meter"), (8, "Calibration Date:", "='Cal-Info'!$B$3"),
        (9, "Report ID:", "='Cal-Info'!$B$4"),
        (10, "Legislation:", "§ 1065.307 Linearity verification. - [79 FR 23766, Apr. 28, 2014]"),
        (12, "Reference Device Data", None), (13, "Description:", "Fluke 8588A DVM digital voltage meter"),
        (14, "Manufacturer:", "Fluke"), (15, "Serial Number:", "PM" + str(info["DVM Gage No"])),
        (16, "Calibration Date:", "='Cal-Info'!$B$14"), (17, "Instrument Range:", "100 mV to 1000 V"),
    ]
    for r, a, b in meta:
        cc = ws.cell(r, 1, a)
        cc.font = BOLD if b is None else NORM
        if b is None:
            cc.fill = SECTION
        if b is not None:
            ws.cell(r, 2, b).font = NORM
    c1 = ws.cell(19, 1, "Test Results")
    c1.font = BOLD
    c1.fill = SECTION
    c2 = ws.cell(19, 7, "Limits")
    c2.font = BOLD
    c2.fill = SECTION
    ws.cell(20, 1, "Description").font = BOLD
    ws.cell(20, 4, "Calculated").font = BOLD
    ws.cell(20, 5, "Limits").font = BOLD
    ws.cell(20, 7, "Min").font = BOLD
    ws.cell(20, 8, "Max").font = BOLD
    d1 = DATA_START_ROW
    d2 = d1 - 1 + npts
    xr = f"B{d1}:B{d2}"
    yr = f"C{d1}:C{d2}"
    gr = f"G{d1}:G{d2}"
    SN = f"'{sheetname}'"
    ws.cell(21, 1, "a1 Slope").font = NORM
    ws.cell(21, 4).value = f"=SLOPE({yr},{xr})"
    ws.cell(21, 4).number_format = "0.0000000"
    ws.cell(21, 5).value = '=G21&" ≤ a1 ≤ "&H21'
    ws.cell(21, 7).value = f"=1-VLOOKUP({SN}!B3,{REF_RANGE},4,FALSE)"
    ws.cell(21, 7).number_format = "0.00"
    ws.cell(21, 8).value = f"=1+VLOOKUP({SN}!B3,{REF_RANGE},4,FALSE)"
    ws.cell(21, 8).number_format = "0.00"
    ws.cell(22, 1, "|Xmin(a1-1)+a0| Interception").font = NORM
    ws.cell(22, 4).value = f"=ABS(INDEX({xr},MATCH(MIN({gr}),{gr},0),1)*(D21-1)+INTERCEPT({yr},{xr}))"
    ws.cell(22, 4).number_format = "0.0000000"
    ws.cell(22, 5).value = '="≤ "&H22'
    ws.cell(22, 8).value = f"=VLOOKUP({SN}!B3,{REF_RANGE},3,FALSE)*MAX({gr})"
    ws.cell(22, 8).number_format = "0.000"
    ws.cell(23, 1, "SEE Standard Estimation of Error").font = NORM
    ws.cell(23, 4).value = f"=STEYX({yr},{xr})"
    ws.cell(23, 4).number_format = "0.0000000"
    ws.cell(23, 5).value = '="≤ "&H23'
    ws.cell(23, 8).value = f"=VLOOKUP({SN}!B3,{REF_RANGE},5,FALSE)*MAX({gr})"
    ws.cell(23, 8).number_format = "0.000"
    ws.cell(24, 1, "r2 Coefficient").font = NORM
    ws.cell(24, 4).value = f"=CORREL({yr},{xr})^2"
    ws.cell(24, 4).number_format = "0.0000000"
    ws.cell(24, 5).value = '="≥ "&G24'
    ws.cell(24, 7).value = f"=VLOOKUP({SN}!B3,{REF_RANGE},6,FALSE)"
    ws.cell(24, 7).number_format = "0.00"
    ws.cell(25, 1, "Test State").font = BOLD
    ws.cell(25, 4).value = (
        f'=IF(AND(COUNT(A{d1}:A{d2})>=3,D21>=G21,D21<=H21,D22<=H22,D23<=H23,D24>=G24),"PASSED","FAILED")'
    )
    ws.cell(25, 4).font = BOLD
    _result_cf(ws, "D25")
    for cc, txt in [(1, "Index"), (2, "Reference [V]"), (3, "Sample [V]"), (4, "Dev. [V]"), (5, "Dev. [%]"), (7, "Abs(ref)")]:
        h = ws.cell(27, cc, txt)
        h.font = BOLD
        h.fill = TBLHDR
    for i in range(npts):
        r = d1 + i
        ws.cell(r, 1).value = (f'=IF(B{r}="","",1)' if i == 0 else f'=IF(B{r}="","",A{r - 1}+1)')
        bc = ws.cell(r, 2)
        bc.value = vsc[i][0]
        bc.fill = YEL
        bc.number_format = "0.000000"
        cc2 = ws.cell(r, 3)
        cc2.value = vsc[i][1]
        cc2.fill = YEL
        cc2.number_format = "0.000000"
        ws.cell(r, 4).value = f'=IF(A{r}="","",C{r}-B{r})'
        ws.cell(r, 4).number_format = "0.0000000"
        ws.cell(r, 5).value = f'=IF(D{r}="","",IF(B{r}=0,"",D{r}/B{r}*100))'
        ws.cell(r, 5).number_format = "0.0000"
        ws.cell(r, 7).value = f'=IF(B{r}="","",ABS(B{r}))'
        ws.cell(r, 7).number_format = "0.00"
        for cc in range(1, 8):
            ws.cell(r, cc).font = NORM
    ws.cell(d2 + 3, 1, "Notes: Linearity Check Version 3.0 (recomputed)").font = NORM
    ws.cell(d2 + 4, 1, "GPN: 1* | Document No: D01* | SAP Material No: TNA99M*").font = NORM
    for col, w in {"A": 19.0, "B": 21.1, "C": 23.6, "D": 25.9, "E": 21.7, "G": 22.1, "H": 23.6}.items():
        ws.column_dimensions[col].width = w


def _sheet_channels(wb, data: CalibrationData, info: dict):
    for spec in channel_specs():
        if spec.range_key == "r10":
            L = gcl(12 + spec.channel)
            vsc = [(v, f"='{PROTO}'!{L}{18 + i}") for i, v in enumerate(data.r10.voltages)]
        elif spec.range_key == "r20":
            vsc = [(v, f"='{PROTO}'!N{30 + i}") for i, v in enumerate(data.r20.voltages)]
        else:
            vsc = [(v, f"='{PROTO}'!N{41 + i}") for i, v in enumerate(data.r75.voltages)]
        _make_channel(wb, spec.sheet, spec.range_label, spec.npts(data), vsc, info)


def _sheet_summary(wb, info: dict):
    ws = wb.create_sheet("Summary", 1)
    ws.cell(1, 1, "Linearity Verification – Summary").font = Font(name="Calibri", bold=True, size=20)
    ws.cell(2, 1, f"AVL X-Meter  |  Serial No {info['Serial No']}  |  Report {info['Report ID']}").font = NORM
    ws.cell(3, 1, f"Calibration Date {info['Calibration Date']}  |  per § 1065.307").font = NORM
    hdr = ["Channel", "Range", "a1 (slope)", "a1 limits", "|Interc.|", "Interc. max", "SEE", "SEE max", "r²", "r² min", "Result"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(5, j, h)
        c.font = BOLD
        c.fill = TBLHDR
    chans = [(f"±10V CH{ch}", f"CH{ch}", "±10V") for ch in range(1, 9)] + [
        ("±20V CH9", "CH9", "±20V"), ("±75V CH10", "CH10", "±75V"),
    ]
    for i, (sn, cl, rg) in enumerate(chans):
        r = 6 + i
        q = f"'{sn}'"
        ws.cell(r, 1, cl).font = NORM
        ws.cell(r, 2, rg).font = NORM
        for col, ref, fmt in [
            (3, "D21", "0.0000000"), (4, "E21", None), (5, "D22", "0.0000000"), (6, "H22", "0.000"),
            (7, "D23", "0.0000000"), (8, "H23", "0.000"), (9, "D24", "0.0000000"), (10, "G24", "0.00"),
        ]:
            ws.cell(r, col).value = f"={q}!{ref}"
            ws.cell(r, col).font = NORM
            if fmt:
                ws.cell(r, col).number_format = fmt
        ws.cell(r, 11).value = f"={q}!D25"
        ws.cell(r, 11).font = BOLD
        _result_cf(ws, f"K{r}")
    ws.column_dimensions["A"].width = 10
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["K"].width = 11


def build_workbook(data: CalibrationData, out_path: Path) -> Path:
    """Build the full 14-sheet Linearity Verification workbook and save it."""
    info = data.info
    wb = openpyxl.Workbook()
    _sheet_cal_info(wb, info)
    _sheet_references(wb)
    _sheet_protocol(wb, data)
    _sheet_channels(wb, data, info)
    _sheet_summary(wb, info)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
